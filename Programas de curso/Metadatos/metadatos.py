# -*- coding: utf-8 -*-
"""
Genera un archivo .xlsx por cada elemento del JSON usando Metadata.xlsx como plantilla.

Instalar dependencias:
    py -m pip install openpyxl
"""

import json
import re
from pathlib import Path
from openpyxl import load_workbook

# Intentar importar soporte de rich text (openpyxl >= 3.1)
try:
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    HAS_RICHTEXT = True
except ImportError:
    HAS_RICHTEXT = False

# ----------------------------------------------------------------------
# Configuración
# ----------------------------------------------------------------------
URL_PLANTILLA_XLSX = "D:\\Micrositios\\cenfotec\\Programas de curso\\Metadatos\\Metadata.xlsx"

TEMPLATE = Path(r"D:\\Micrositios\\cenfotec\\Programas de curso\\Metadatos\\Metadata.xlsx")
JSON_PATH = Path(r"D:\\Micrositios\\cenfotec\\Programas de curso\\Metadatos\\cursos.json")
OUT_DIR   = Path(r"D:\\Micrositios\\cenfotec\\Programas de curso\\Metadatos\Salida")

# Mapeo: clave dentro de las llaves  ->  campo del JSON
FIELD_MAP = {
    "Código y nombre del curso":          "codigo_y_nombre",
    "descripción":                        "descripcion",
    "deswcripción":                       "descripcion",   # por si la plantilla tiene el typo
    "descripcion":                        "descripcion",
    "habilidades":                        "habilidades",
    "habilidad 1":                        "habilidad_1",
    "habilidad 2":                        "habilidad_2",
    "habilidad 3":                        "habilidad_3",
    "nivel sfia":                         "nivel_sfia",
    "habilidad / Área de conocimiento":   "habilidad_area_conocimiento",
    "link sfia":                          "link_sfia",
}

# Detecta {{{ ... }}} ignorando espacios alrededor del contenido
PLACEHOLDER_RE = re.compile(r"\{\{\{\s*(.*?)\s*\}\}\}", re.DOTALL)


# ----------------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------------
def apply_replacements(text: str, item: dict) -> str:
    """Sustituye todos los {{{...}}} conocidos por los valores del item."""
    def _repl(m):
        key = m.group(1).strip()
        field = FIELD_MAP.get(key)
        if field is None:
            return m.group(0)              # desconocido -> se deja tal cual
        value = item.get(field, "")
        return "" if value is None else str(value)
    return PLACEHOLDER_RE.sub(_repl, text)


def flatten(value) -> str | None:
    """Devuelve el texto plano de una celda (str o CellRichText)."""
    if isinstance(value, str):
        return value
    if HAS_RICHTEXT and isinstance(value, CellRichText):
        return "".join(r if isinstance(r, str) else r.text for r in value)
    return None


def replace_in_cell(cell, item: dict):
    """Reemplaza los placeholders dentro de una celda (soporta rich text)."""
    v = cell.value
    if v is None:
        return

    flat = flatten(v)
    if flat is None or "{{{" not in flat:
        return

    new_flat = apply_replacements(flat, item)

    # Celda simple
    if isinstance(v, str):
        if new_flat != v:
            cell.value = new_flat
        return

    # Celda con rich text: intentar reemplazo por run
    if HAS_RICHTEXT and isinstance(v, CellRichText):
        changed = False
        new_runs = []
        for r in v:
            if isinstance(r, str):
                nr = apply_replacements(r, item)
                if nr != r:
                    changed = True
                new_runs.append(nr)
            else:  # TextBlock
                nr = apply_replacements(r.text, item)
                if nr != r.text:
                    changed = True
                    new_runs.append(TextBlock(r.font, nr))
                else:
                    new_runs.append(r)

        joined = "".join(r if isinstance(r, str) else r.text for r in new_runs)

        # Si aún quedan {{{..., probablemente están partidos entre runs -> aplanar
        if "{{{" in joined:
            cell.value = new_flat
        elif changed:
            cell.value = CellRichText(*new_runs)


def sanitize_filename(name: str) -> str:
    """Quita caracteres inválidos para nombres de archivo en Windows."""
    name = re.sub(r'[\\/:*?"<>|]', "_", str(name))
    name = re.sub(r"\s+", " ", name).strip()
    return name or "curso"


# ----------------------------------------------------------------------
# Programa principal
# ----------------------------------------------------------------------
def main():
    data = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    if not isinstance(data, list):
        raise SystemExit("El JSON debe ser una lista de cursos.")

    OUT_DIR.mkdir(parents=True, exist_ok=True)

    for idx, item in enumerate(data, start=1):
        wb = load_workbook(TEMPLATE)

        # Recorre TODAS las hojas y celdas
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    replace_in_cell(cell, item)

        base = item.get("codigo_y_nombre") or f"curso_{idx}"
        out_path = OUT_DIR / f"Metadata {sanitize_filename(base)}.xlsx"
        wb.save(out_path)
        print(f"[{idx}/{len(data)}] Generado: {out_path}")


if __name__ == "__main__":
    main()

# Para correr el programa usar python '.\parser de cursos.py'

# py -m pip install openpyxl