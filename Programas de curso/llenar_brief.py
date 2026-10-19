#!/usr/bin/env python3
"""
llenar_brief.py
Copia la información de un documento de respuestas (p. ej. "mercadeo.docx")
a la plantilla del brief (p. ej. "Brief_mercadeo.docx"), reemplazando cada
marcador entre corchetes ([valor], [duración], ...) por el contenido de la
sección del mismo nombre en el documento de respuestas.

Uso:
    python llenar_brief.py mercadeo.docx Brief_mercadeo.docx Brief_lleno.docx
    python llenar_brief.py mercadeo.docx Brief_mercadeo.docx Brief_lleno.docx --quitar-marcas

Opciones:
    --quitar-marcas   Elimina las marcas "[Información tomada del archivo ...]"
                      y "[Información propuesta: revisar]". Las marcas
                      "[Pendiente: ...]" se conservan siempre.

Cómo empareja: busca en ambos documentos los títulos de sección (se comparan
sin mayúsculas, tildes ni signos). Todo lo que hay en el documento de
respuestas entre un título y el siguiente se copia en lugar del marcador que
sigue a ese título en el brief. Las secciones del brief que ya tienen texto
(sin marcador) no se tocan.

Requiere: pip install python-docx
"""
import argparse
import copy
import re
import sys
import unicodedata

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

MARCADOR = re.compile(r"^\s*\[[^\[\]]+\]\s*$")
MARCAS_REVISION = re.compile(r"\s*\[Información (tomada del archivo original[^\]]*|propuesta: revisar)\]")
# ": [Información ...]" seguido de más texto = nota de procedencia, no contenido
NOTA_ATRIBUCION = re.compile(r":\s*\[Información [^\]]*\]\s*\S")


# ----------------------------------------------------------------- utilidades
def normalizar(texto: str) -> str:
    """Minúsculas, sin tildes ni signos, espacios simples."""
    t = unicodedata.normalize("NFD", texto.lower())
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    t = re.sub(r"[^a-z0-9ñ ]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def texto_parrafo(p_el) -> str:
    """Texto de un párrafo incluyendo el de los hipervínculos."""
    return "".join(t.text or "" for t in p_el.iter(qn("w:t")))


def elementos_cuerpo(documento):
    """Párrafos y tablas del cuerpo, en orden."""
    for el in documento.element.body.iterchildren():
        if el.tag in (qn("w:p"), qn("w:tbl")):
            yield el


def es_titulo_estilo(p_el, documento) -> bool:
    estilo = Paragraph(p_el, documento).style
    return estilo is not None and estilo.name.lower().startswith(("heading", "título", "titulo"))


# ------------------------------------------------------- lectura de la plantilla
def titulos_y_marcadores(brief):
    """Devuelve (titulos, marcadores).
    titulos: conjunto de claves normalizadas de todos los títulos del brief.
    marcadores: lista de (clave_titulo, texto_titulo, elemento_marcador)."""
    elems = list(elementos_cuerpo(brief))
    titulos = {}
    marcadores = []
    ultimo = None
    for i, el in enumerate(elems):
        if el.tag != qn("w:p"):
            continue
        txt = texto_parrafo(el).strip()
        if not txt:
            continue
        siguiente = elems[i + 1] if i + 1 < len(elems) else None
        siguiente_es_marcador = (
            siguiente is not None
            and siguiente.tag == qn("w:p")
            and MARCADOR.match(texto_parrafo(siguiente))
        )
        if MARCADOR.match(txt):
            if ultimo:
                marcadores.append((ultimo[0], ultimo[1], el))
            continue
        if es_titulo_estilo(el, brief) or siguiente_es_marcador:
            ultimo = (normalizar(txt), txt)
            titulos[ultimo[0]] = txt
    return titulos, marcadores


# ------------------------------------------------------- lectura de respuestas
def secciones_respuestas(fuente, claves_titulos):
    """Divide el documento de respuestas en secciones según los títulos del brief."""
    secciones = {}
    actual = None
    for el in elementos_cuerpo(fuente):
        if el.tag == qn("w:p"):
            clave = normalizar(texto_parrafo(el))
            if clave in claves_titulos:
                actual = clave
                secciones.setdefault(actual, [])
                continue
        if actual is not None:
            if el.tag == qn("w:p") and not texto_parrafo(el).strip():
                continue
            secciones[actual].append(el)
    return secciones


# --------------------------------------------------------------- construcción
class Constructor:
    """Crea en el brief copias con su propio formato de párrafos y tablas."""

    def __init__(self, brief, fuente, quitar_marcas):
        self.brief = brief
        self.fuente = fuente
        self.quitar_marcas = quitar_marcas
        self.num_vinetas = self._buscar_vinetas()
        tablas = brief.tables
        self.tabla_modelo = tablas[0]._tbl if tablas else None

    # -- viñetas
    def _buscar_vinetas(self):
        """numId de una lista con viñetas en la plantilla (o crea una)."""
        numbering = self.brief.part.numbering_part.element
        abstract_bullet = None
        for an in numbering.findall(qn("w:abstractNum")):
            lvl0 = an.find(qn("w:lvl"))
            fmt = lvl0.find(qn("w:numFmt")) if lvl0 is not None else None
            if fmt is not None and fmt.get(qn("w:val")) == "bullet":
                abstract_bullet = an.get(qn("w:abstractNumId"))
                break
        if abstract_bullet is None:
            return None
        nums = numbering.findall(qn("w:num"))
        nuevo_id = str(max(int(n.get(qn("w:numId"))) for n in nums) + 1 if nums else 1)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), nuevo_id)
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), abstract_bullet)
        num.append(ref)
        numbering.append(num)
        return nuevo_id

    # -- runs
    @staticmethod
    def _rpr_base(p_modelo):
        r = p_modelo.find(qn("w:r"))
        rpr = r.find(qn("w:rPr")) if r is not None else None
        return copy.deepcopy(rpr) if rpr is not None else OxmlElement("w:rPr")

    @staticmethod
    def _activar(rpr, etiqueta, valor=None):
        el = rpr.find(qn(etiqueta))
        if el is None:
            el = OxmlElement(etiqueta)
            rpr.append(el)
        if valor is not None:
            el.set(qn("w:val"), valor)

    def _nuevo_run(self, r_fuente, rpr_base, es_link=False):
        r = OxmlElement("w:r")
        rpr = copy.deepcopy(rpr_base)
        src = r_fuente.find(qn("w:rPr"))
        if src is not None:
            for tag in ("w:b", "w:i", "w:u", "w:strike"):
                el = src.find(qn(tag))
                if el is not None and el.get(qn("w:val")) not in ("0", "false"):
                    self._activar(rpr, tag, el.get(qn("w:val")))
        if es_link:
            self._activar(rpr, "w:color", "0563C1")
            self._activar(rpr, "w:u", "single")
        if len(rpr):
            r.append(rpr)
        for hijo in r_fuente:
            if hijo.tag == qn("w:t"):
                t = OxmlElement("w:t")
                t.text = hijo.text
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                r.append(t)
            elif hijo.tag in (qn("w:tab"), qn("w:br")):
                r.append(OxmlElement("w:" + hijo.tag.split("}")[1]))
        return r

    def _limpiar_marcas(self, p):
        """Quita las marcas de revisión dentro de los w:t del párrafo."""
        for t in p.iter(qn("w:t")):
            if t.text:
                t.text = MARCAS_REVISION.sub("", t.text)
        # quitar espacios sobrantes al final del párrafo
        ts = [t for t in p.iter(qn("w:t")) if t.text]
        if ts:
            ts[-1].text = ts[-1].text.rstrip()
        for r in list(p.iter(qn("w:r"))):
            if not "".join(t.text or "" for t in r.findall(qn("w:t"))) and r.find(qn("w:t")) is not None:
                r.getparent().remove(r)

    def parrafo(self, p_fuente, p_modelo):
        p = OxmlElement("w:p")
        ppr_modelo = p_modelo.find(qn("w:pPr"))
        ppr = copy.deepcopy(ppr_modelo) if ppr_modelo is not None else OxmlElement("w:pPr")
        viejo = ppr.find(qn("w:numPr"))
        if viejo is not None:
            ppr.remove(viejo)
        ppr_src = p_fuente.find(qn("w:pPr"))
        es_lista = ppr_src is not None and ppr_src.find(qn("w:numPr")) is not None
        if es_lista and self.num_vinetas:
            numpr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), "0")
            numid = OxmlElement("w:numId")
            numid.set(qn("w:val"), self.num_vinetas)
            numpr.append(ilvl)
            numpr.append(numid)
            pstyle = ppr.find(qn("w:pStyle"))
            ppr.insert(1 if pstyle is not None else 0, numpr)
        p.append(ppr)

        rpr_base = self._rpr_base(p_modelo)
        for b in ("w:b", "w:i", "w:u"):  # el formato real lo da la fuente
            el = rpr_base.find(qn(b))
            if el is not None:
                rpr_base.remove(el)
        if es_lista and not self.num_vinetas:
            r = OxmlElement("w:r")
            r.append(copy.deepcopy(rpr_base))
            t = OxmlElement("w:t")
            t.text = "•  "
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            r.append(t)
            p.append(r)

        for hijo in p_fuente:
            if hijo.tag == qn("w:r"):
                p.append(self._nuevo_run(hijo, rpr_base))
            elif hijo.tag == qn("w:hyperlink"):
                rid = hijo.get(qn("r:id"))
                url = None
                if rid and rid in self.fuente.part.rels:
                    url = self.fuente.part.rels[rid].target_ref
                contenedor = p
                if url:
                    nuevo_rid = self.brief.part.relate_to(url, RT.HYPERLINK, is_external=True)
                    h = OxmlElement("w:hyperlink")
                    h.set(qn("r:id"), nuevo_rid)
                    h.set(qn("w:history"), "1")
                    p.append(h)
                    contenedor = h
                for r in hijo.iter(qn("w:r")):
                    contenedor.append(self._nuevo_run(r, rpr_base, es_link=bool(url)))
        if self.quitar_marcas:
            texto = texto_parrafo(p)
            # Notas que solo explican de dónde salió un dato (p. ej.
            # "Datos de la tabla: [marca], salvo ...: [marca].") se omiten.
            if NOTA_ATRIBUCION.search(texto):
                return None
            self._limpiar_marcas(p)
        return p

    # -- tablas
    def tabla(self, tbl_fuente, p_modelo):
        t_src = Table(tbl_fuente, self.fuente)
        filas = [[c for c in fila.cells] for fila in t_src.rows]
        ncol = max(len(f) for f in filas)
        if self.tabla_modelo is None:
            return copy.deepcopy(tbl_fuente)

        modelo = self.tabla_modelo
        filas_modelo = modelo.findall(qn("w:tr"))
        tc_enc = filas_modelo[0].find(qn("w:tc"))
        tc_cuerpo = (filas_modelo[1] if len(filas_modelo) > 1 else filas_modelo[0]).find(qn("w:tc"))
        tblpr = copy.deepcopy(modelo.find(qn("w:tblPr")))
        ancho_total = int(tblpr.find(qn("w:tblW")).get(qn("w:w")) or 8800)

        # anchos proporcionales a la tabla original
        grid_src = [int(g.get(qn("w:w")) or 1) for g in tbl_fuente.find(qn("w:tblGrid")).findall(qn("w:gridCol"))]
        if len(grid_src) != ncol or sum(grid_src) == 0:
            grid_src = [1] * ncol
        anchos = [round(ancho_total * g / sum(grid_src)) for g in grid_src]

        tbl = OxmlElement("w:tbl")
        tbl.append(tblpr)
        grid = OxmlElement("w:tblGrid")
        for a in anchos:
            gc = OxmlElement("w:gridCol")
            gc.set(qn("w:w"), str(a))
            grid.append(gc)
        tbl.append(grid)

        for i, fila in enumerate(filas):
            tr = OxmlElement("w:tr")
            if i == 0:
                trpr = filas_modelo[0].find(qn("w:trPr"))
                if trpr is not None:
                    tr.append(copy.deepcopy(trpr))
            tc_mod = tc_enc if i == 0 else tc_cuerpo
            for j, celda in enumerate(fila):
                tc = OxmlElement("w:tc")
                tcpr = copy.deepcopy(tc_mod.find(qn("w:tcPr")))
                tcpr.find(qn("w:tcW")).set(qn("w:w"), str(anchos[j]))
                tc.append(tcpr)
                p_mod = tc_mod.find(qn("w:p"))
                for p_src in celda._tc.findall(qn("w:p")):
                    p_new = self.parrafo(p_src, p_mod)
                    if p_new is None:
                        continue
                    jc = OxmlElement("w:jc")  # celdas alineadas a la izquierda
                    jc.set(qn("w:val"), "left")
                    p_new.find(qn("w:pPr")).append(jc)
                    if i == 0:  # encabezado: usar el formato del encabezado modelo
                        rpr_enc = self._rpr_base(p_mod)
                        for r in p_new.findall(qn("w:r")):
                            viejo = r.find(qn("w:rPr"))
                            if viejo is not None:
                                r.remove(viejo)
                            r.insert(0, copy.deepcopy(rpr_enc))
                            for t in r.findall(qn("w:t")):
                                t.text = (t.text or "").upper()
                    tc.append(p_new)
                tr.append(tc)
            tbl.append(tr)
        return tbl


# ---------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Llena el brief con la información del documento de respuestas.")
    ap.add_argument("respuestas", help="Documento con la información (p. ej. mercadeo.docx)")
    ap.add_argument("plantilla", help="Brief con los marcadores [ ] (p. ej. Brief_mercadeo.docx)")
    ap.add_argument("salida", help="Archivo de salida (.docx)")
    ap.add_argument("--quitar-marcas", action="store_true",
                    help="Quitar [Información tomada...] y [Información propuesta: revisar]")
    args = ap.parse_args()

    fuente = docx.Document(args.respuestas)
    brief = docx.Document(args.plantilla)

    titulos, marcadores = titulos_y_marcadores(brief)
    secciones = secciones_respuestas(fuente, set(titulos))
    constructor = Constructor(brief, fuente, args.quitar_marcas)

    llenadas, faltantes = [], []
    usadas = set()
    for clave, titulo, marcador in marcadores:
        contenido = secciones.get(clave)
        if not contenido:
            faltantes.append((titulo, texto_parrafo(marcador).strip()))
            continue
        ancla = marcador
        for el in contenido:
            nuevo = (constructor.parrafo(el, marcador) if el.tag == qn("w:p")
                     else constructor.tabla(el, marcador))
            if nuevo is None:
                continue
            ancla.addnext(nuevo)
            ancla = nuevo
        marcador.getparent().remove(marcador)
        llenadas.append(titulo)
        usadas.add(clave)

    brief.save(args.salida)

    print(f"Listo: {args.salida}")
    print(f"\nSecciones llenadas ({len(llenadas)}):")
    for t in llenadas:
        print(f"  ✓ {t}")
    if faltantes:
        print(f"\nMarcadores sin información en '{args.respuestas}' ({len(faltantes)}):")
        for t, m in faltantes:
            print(f"  ✗ {t}  →  {m}")
    sin_usar = [titulos[k] for k, v in secciones.items() if k not in usadas and v]
    if sin_usar:
        print("\nSecciones con contenido en las respuestas que no se copiaron "
              "(en el brief ya tienen texto o no tienen marcador):")
        for t in sin_usar:
            print(f"  · {t}")


if __name__ == "__main__":
    sys.exit(main())
